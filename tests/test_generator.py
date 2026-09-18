import json
import os
from pathlib import Path
from unittest.mock import patch, MagicMock

import docx
from docx.shared import RGBColor
import pytest

from repertorio.generator import generate_repertoire


MOCK_SOLR_RESPONSE = {
    "response": {
        "numFound": 1,
        "docs": [
            {
                "art": "Soda Stereo",
                "dns": "soda-stereo",
                "txt": "De Música Ligera",
                "url": "de-musica-ligera",
            }
        ],
    }
}

MOCK_HTML_PAGE = """<!DOCTYPE html>
<html>
<body>
<h1>De Música Ligera</h1>
<pre class="_crVx" data-chord-content="true">
<div class="kvMV">[Intro] <b>Bm</b>  <b>G</b>  <b>D</b>  <b>A</b>
E|----------------------------|
B|----------------------------|
G|----------------------------|
D|----------------------------|
A|----------------------------|
E|----------------------------|
</div>
<div class="kvMV"><b>Bm</b>        <b>G</b>  <b>D</b>
Ella durmio
      <b>A</b>           <b>Bm</b>   <b>G</b>  <b>D</b>  <b>A</b>
Al calor de las masas
</div>
</pre>
</body>
</html>
"""


@pytest.fixture
def temp_workspace(tmp_path):
    songs_file = tmp_path / "canciones.txt"
    songs_file.write_text("de musica ligera\n", encoding="utf-8")
    output_docx = tmp_path / "repertorio.docx"
    cache_dir = tmp_path / ".cache_cifras"
    return {
        "songs_file": songs_file,
        "output_docx": output_docx,
        "cache_dir": cache_dir,
    }


def test_generator_end_to_end(temp_workspace):
    songs_file = temp_workspace["songs_file"]
    output_docx = temp_workspace["output_docx"]
    cache_dir = temp_workspace["cache_dir"]

    # Mock HTTP calls
    with patch("urllib.request.urlopen") as mock_urlopen:
        mock_solr_resp = MagicMock()
        mock_solr_resp.read.return_value = json.dumps(MOCK_SOLR_RESPONSE).encode("utf-8")
        mock_solr_resp.__enter__.return_value = mock_solr_resp

        mock_page_resp = MagicMock()
        mock_page_resp.read.return_value = MOCK_HTML_PAGE.encode("utf-8")
        mock_page_resp.__enter__.return_value = mock_page_resp

        mock_urlopen.side_effect = [mock_solr_resp, mock_page_resp]

        # 1. First execution: should query network and create cache
        stats = generate_repertoire(songs_file, output_docx, cache_dir=cache_dir)

        assert stats["total"] == 1
        assert stats["downloaded"] == 1
        assert stats["cached"] == 0
        assert stats["failed"] == 0
        assert mock_urlopen.call_count == 2

    # Verify output file exists
    assert output_docx.exists()

    # Verify cache files: 1 in songs/ and 1 index.json
    songs_dir = cache_dir / "songs"
    cache_files = list(songs_dir.glob("*.json"))
    assert len(cache_files) == 1
    with open(cache_files[0], encoding="utf-8") as f:
        cached_data = json.load(f)
    assert cached_data["artist"] == "Soda Stereo"
    assert cached_data["title"] == "De Música Ligera"

    index_file = cache_dir / "index.json"
    assert index_file.exists()
    with open(index_file, encoding="utf-8") as f:
        index_data = json.load(f)
    assert index_data.get("de musica ligera") == "soda-stereo_de-musica-ligera"

    # Verify DOCX structure
    doc = docx.Document(output_docx)
    section = doc.sections[0]

    # 2 columns configuration check
    cols = section._sectPr.xpath("./w:cols")
    assert len(cols) > 0
    assert cols[0].get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}num") == "2"

    all_text = "\n".join([p.text for p in doc.paragraphs])
    # Should include title and lyrics
    assert "De Música Ligera" in all_text
    assert "Soda Stereo" in all_text
    assert "Ella durmio" in all_text

    # Tablature lines must be filtered out
    assert "E|----------------------------|" not in all_text
    assert "A|----------------------------|" not in all_text

    # Verify chord styling (Consolas, bold, orange RGB(230, 81, 0))
    chord_runs = []
    for p in doc.paragraphs:
        for r in p.runs:
            if r.bold and r.font.color and r.font.color.rgb == RGBColor(230, 81, 0):
                chord_runs.append(r.text)

    assert "Bm" in chord_runs
    assert "G" in chord_runs
    assert "D" in chord_runs
    assert "A" in chord_runs

    # 2. Second execution: should use cache, zero network calls
    with patch("urllib.request.urlopen") as mock_urlopen_2:
        stats_2 = generate_repertoire(songs_file, output_docx, cache_dir=cache_dir)
        assert stats_2["total"] == 1
        assert stats_2["downloaded"] == 0
        assert stats_2["cached"] == 1
        assert stats_2["failed"] == 0
        assert mock_urlopen_2.call_count == 0


def test_generate_from_songs(temp_workspace):
    from repertorio.generator import generate_from_songs

    output_docx = temp_workspace["output_docx"]
    cache_dir = temp_workspace["cache_dir"]
    songs = [
        {"artist": "Soda Stereo", "title": "De Música Ligera", "dns": "soda-stereo", "url": "de-musica-ligera"}
    ]

    with patch("urllib.request.urlopen") as mock_urlopen:
        mock_page_resp = MagicMock()
        mock_page_resp.read.return_value = MOCK_HTML_PAGE.encode("utf-8")
        mock_page_resp.__enter__.return_value = mock_page_resp
        mock_urlopen.return_value = mock_page_resp

        progress_calls = []
        def on_progress(cur, tot, msg):
            progress_calls.append((cur, tot, msg))

        stats = generate_from_songs(songs, output_docx, cache_dir=cache_dir, progress_callback=on_progress)
        assert stats["total"] == 1
        assert stats["downloaded"] == 1
        assert stats["failed"] == 0
        assert len(progress_calls) >= 2
        assert output_docx.exists()


def test_generate_from_songs_with_transposition(temp_workspace):
    from repertorio.generator import generate_from_songs

    output_docx = temp_workspace["output_docx"]
    cache_dir = temp_workspace["cache_dir"]
    songs = [
        {
            "artist": "Soda Stereo",
            "title": "De Música Ligera",
            "dns": "soda-stereo",
            "url": "de-musica-ligera",
            "key": "Bm",
            "semitones": 2,
        }
    ]

    with patch("urllib.request.urlopen") as mock_urlopen:
        mock_page_resp = MagicMock()
        mock_page_resp.read.return_value = MOCK_HTML_PAGE.encode("utf-8")
        mock_page_resp.__enter__.return_value = mock_page_resp
        mock_urlopen.return_value = mock_page_resp

        stats = generate_from_songs(songs, output_docx, cache_dir=cache_dir)
        assert stats["total"] == 1
        assert stats["downloaded"] == 1

        doc = docx.Document(output_docx)
        all_text = "\n".join([p.text for p in doc.paragraphs])
        # Transposed key with original reference
        assert "Tono: C#m (Original: Bm)" in all_text

        # Transposed chords: Bm (+2) -> C#m, G (+2) -> A, D (+2) -> E, A (+2) -> B
        chord_runs = []
        for p in doc.paragraphs:
            for r in p.runs:
                if r.bold and r.font.color and r.font.color.rgb == RGBColor(230, 81, 0):
                    chord_runs.append(r.text)

        assert "C#m" in chord_runs
        assert "A" in chord_runs
        assert "E" in chord_runs
        assert "B" in chord_runs
        # Original chords should not be present
        assert "Bm" not in chord_runs


def test_generate_from_songs_with_override(temp_workspace):
    from repertorio.generator import generate_from_songs
    from repertorio.parser import parse_text_to_lines

    output_docx = temp_workspace["output_docx"]
    cache_dir = temp_workspace["cache_dir"]

    custom_chord_sheet = (
        "[Intro] Em  C  G  D\n"
        "\n"
        "Em        C\n"
        "Custom modified lyrics line\n"
        "      D           Em\n"
        "Another custom arranged line"
    )
    override_lines = parse_text_to_lines(custom_chord_sheet)

    songs = [
        {
            "artist": "Soda Stereo",
            "title": "De Música Ligera",
            "dns": "soda-stereo",
            "url": "de-musica-ligera",
            "key": "Em",
            "semitones": 0,
            "override": {
                "lines": override_lines,
                "is_modified": True,
            },
        }
    ]

    with patch("urllib.request.urlopen") as mock_urlopen:
        stats = generate_from_songs(songs, output_docx, cache_dir=cache_dir)
        # Should not make any network requests
        assert mock_urlopen.call_count == 0
        assert stats["total"] == 1
        assert stats["downloaded"] == 0
        assert stats["cached"] == 1
        assert stats["failed"] == 0

    assert output_docx.exists()
    doc = docx.Document(output_docx)
    all_text = "\n".join([p.text for p in doc.paragraphs])

    # Should contain custom lyrics from override, NOT CifraClub original
    assert "Custom modified lyrics line" in all_text
    assert "Another custom arranged line" in all_text
    assert "Ella durmio" not in all_text

    # Verify custom chords are styled with bold orange
    chord_runs = []
    for p in doc.paragraphs:
        for r in p.runs:
            if r.bold and r.font.color and r.font.color.rgb == RGBColor(230, 81, 0):
                chord_runs.append(r.text)

    assert "Em" in chord_runs
    assert "C" in chord_runs
    assert "G" in chord_runs
    assert "D" in chord_runs

    # Global cache should remain untouched (no song files written)
    songs_dir = cache_dir / "songs"
    assert not songs_dir.exists() or len(list(songs_dir.glob("*.json"))) == 0


def test_generate_from_songs_with_override_and_transposition(temp_workspace):
    from repertorio.generator import generate_from_songs
    from repertorio.parser import parse_text_to_lines

    output_docx = temp_workspace["output_docx"]
    cache_dir = temp_workspace["cache_dir"]

    custom_chord_sheet = (
        "G        C\n"
        "Acoustic version lyrics\n"
        "      D           G\n"
        "Closing line"
    )
    override_lines = parse_text_to_lines(custom_chord_sheet)

    songs = [
        {
            "artist": "Sample Artist",
            "title": "Sample Song",
            "dns": "sample-artist",
            "url": "sample-song",
            "key": "G",
            "semitones": 2,  # Transpose +2 semitones: G -> A, C -> D, D -> E
            "override": {
                "lines": override_lines,
                "is_modified": True,
            },
        }
    ]

    with patch("urllib.request.urlopen") as mock_urlopen:
        stats = generate_from_songs(songs, output_docx, cache_dir=cache_dir)
        assert mock_urlopen.call_count == 0
        assert stats["total"] == 1
        assert stats["cached"] == 1

    doc = docx.Document(output_docx)
    all_text = "\n".join([p.text for p in doc.paragraphs])

    # Key metadata reflects active transposition and original reference
    assert "Tono: A (Original: G)" in all_text
    assert "Acoustic version lyrics" in all_text

    # Chords must be shifted by +2
    chord_runs = []
    for p in doc.paragraphs:
        for r in p.runs:
            if r.bold and r.font.color and r.font.color.rgb == RGBColor(230, 81, 0):
                chord_runs.append(r.text)

    assert "A" in chord_runs
    assert "D" in chord_runs
    assert "E" in chord_runs
    # Original untransposed chords should not be present
    assert "G" not in chord_runs
    assert "C" not in chord_runs


